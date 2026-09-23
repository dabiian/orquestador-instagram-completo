from __future__ import annotations

import base64
import io
import json
import re
import unicodedata
from datetime import UTC, date, datetime
from typing import Any, Literal
from urllib.parse import urlparse
from uuid import UUID

from pydantic import (
    AnyHttpUrl,
    BaseModel,
    Field,
    ValidationError,
    field_validator,
    model_validator,
)

from orchestrator.application.dtos import FlowSubmitRequest, WordpressPageSetupRequest
from orchestrator.application.import_template import REQUIRED_COLUMNS
from orchestrator.application.use_cases.dispatch_execution import DispatchExecution
from orchestrator.application.use_cases.page_executions import PageExecutionService
from orchestrator.application.use_cases.submit_flow import SubmitFlow
from orchestrator.domain.entities import ExecutionStatus
from orchestrator.domain.ports import RankPositionProvider, SeoAgentManagementRepository

SUPPORT_POST_MODES = {"auto", "state_defined", "legacy_two", "none"}
AdminResource = Literal[
    "campaigns",
    "campaign-services",
    "wordpress-sites",
    "campaign-prompt-rules",
]


class SupportPostsPreference(BaseModel):
    mode: Literal["auto", "state_defined", "legacy_two", "none"] = "auto"


class CampaignPageCreateRequest(BaseModel):
    campaign_id: int = Field(gt=0)
    service_id: int | None = Field(default=None, gt=0)
    wp_page_id: int | None = Field(default=None, gt=0)
    url: AnyHttpUrl
    slug: str | None = Field(default=None, max_length=300)
    page_type: Literal["home", "service", "service_city"] = "service_city"
    language: str = Field(default="es", min_length=2, max_length=10)
    primary_keyword: str = Field(min_length=1, max_length=500)
    secondary_keywords: list[str] = Field(default_factory=list)
    service_scope: Literal["all", "single"] = "all"
    customer_segment: Literal[
        "residential",
        "commercial",
        "both",
        "not_applicable",
        "unknown",
    ] = "unknown"
    city: str | None = Field(default=None, max_length=180)
    state: str | None = Field(default=None, max_length=120)
    country: str = Field(default="US", max_length=10)
    target_location_name: str | None = Field(default=None, max_length=300)
    target_google_maps_url: str | None = Field(default=None, max_length=2000)
    location_type: Literal["city", "neighborhood", "general"] = "city"
    parent_city: str | None = Field(default=None, max_length=180)
    delivery_mode: Literal["remote", "onsite", "hybrid", "unknown"] = "unknown"
    allows_remote: bool | None = None
    allows_onsite: bool | None = None
    physical_visit_required: bool | None = None
    postal_code: str | None = Field(default=None, max_length=30)
    page_template: str = Field(default="elementor_header_footer", max_length=120)
    layout_id: int | None = Field(default=None, gt=0)
    layout_branch: Literal["cities", "services"] | None = None
    parent_campaign_page_id: int | None = Field(default=None, gt=0)
    parent_slug: str | None = Field(default=None, max_length=300)
    parent_required: bool = False
    elementor_form_id: str | None = Field(default=None, max_length=120)
    youtube_channel_id: str | None = Field(
        default=None,
        max_length=120,
        pattern=r"^UC[A-Za-z0-9_-]+$",
    )
    allow_publish: bool = True
    status: Literal["pending", "active", "published"] = "pending"

    @model_validator(mode="after")
    def neighborhood_requires_parent_city(self) -> CampaignPageCreateRequest:
        if self.location_type == "neighborhood" and not self.parent_city:
            raise ValueError("parent_city is required for neighborhood pages")
        return self

    def to_repository_payload(self) -> dict[str, Any]:
        payload = self.model_dump(mode="python")
        payload["url"] = str(self.url)
        payload["slug"] = (self.slug or _slug_from_url(str(self.url))).strip("/")
        payload["secondary_keywords"] = [
            keyword.strip() for keyword in self.secondary_keywords if keyword.strip()
        ]
        if payload.get("parent_slug"):
            payload["parent_slug"] = str(payload["parent_slug"]).strip("/")
        return payload


class CampaignPageUpdateRequest(BaseModel):
    """Partial update of a registered page: only the sent fields are written."""

    service_id: int | None = Field(default=None, gt=0)
    wp_page_id: int | None = Field(default=None, gt=0)
    url: AnyHttpUrl | None = None
    slug: str | None = Field(default=None, max_length=300)
    page_type: Literal["home", "service", "service_city"] | None = None
    language: str | None = Field(default=None, min_length=2, max_length=10)
    primary_keyword: str | None = Field(default=None, min_length=1, max_length=500)
    secondary_keywords: list[str] | None = None
    service_scope: Literal["all", "single"] | None = None
    customer_segment: (
        Literal[
            "residential",
            "commercial",
            "both",
            "not_applicable",
            "unknown",
        ]
        | None
    ) = None
    city: str | None = Field(default=None, max_length=180)
    state: str | None = Field(default=None, max_length=120)
    country: str | None = Field(default=None, max_length=10)
    target_location_name: str | None = Field(default=None, max_length=300)
    target_google_maps_url: str | None = Field(default=None, max_length=2000)
    location_type: Literal["city", "neighborhood", "general"] | None = None
    parent_city: str | None = Field(default=None, max_length=180)
    delivery_mode: Literal["remote", "onsite", "hybrid", "unknown"] | None = None
    allows_remote: bool | None = None
    allows_onsite: bool | None = None
    physical_visit_required: bool | None = None
    postal_code: str | None = Field(default=None, max_length=30)
    page_template: str | None = Field(default=None, max_length=120)
    layout_id: int | None = Field(default=None, gt=0)
    layout_branch: Literal["cities", "services"] | None = None
    parent_campaign_page_id: int | None = Field(default=None, gt=0)
    parent_slug: str | None = Field(default=None, max_length=300)
    parent_required: bool | None = None
    elementor_form_id: str | None = Field(default=None, max_length=120)
    youtube_channel_id: str | None = Field(
        default=None,
        max_length=120,
        pattern=r"^UC[A-Za-z0-9_-]+$",
    )
    allow_publish: bool | None = None
    status: Literal["pending", "active", "published"] | None = None

    def to_repository_payload(self) -> dict[str, Any]:
        payload = self.model_dump(mode="python", exclude_unset=True)
        if not payload:
            raise ValueError("At least one field is required")
        if "url" in payload:
            if self.url is None:
                raise ValueError("url cannot be null")
            payload["url"] = str(self.url)
        if payload.get("slug"):
            payload["slug"] = str(payload["slug"]).strip("/")
        if payload.get("parent_slug"):
            payload["parent_slug"] = str(payload["parent_slug"]).strip("/")
        if payload.get("secondary_keywords") is not None:
            payload["secondary_keywords"] = [
                keyword.strip()
                for keyword in payload["secondary_keywords"]
                if keyword.strip()
            ]
        return payload


class PostGroupCityPayload(BaseModel):
    city: str = Field(min_length=1, max_length=180)
    state: str | None = Field(default=None, max_length=120)


class PostGroupCreateRequest(BaseModel):
    campaign_id: int = Field(gt=0)
    state: str = Field(min_length=2, max_length=120)
    region_name: str = Field(min_length=1, max_length=300)
    group_type: Literal[
        "state_general",
        "county",
        "metro",
        "regional",
        "custom",
    ] = "regional"
    priority: int = Field(default=100, ge=0, le=1000)
    status: Literal["active", "paused", "archived"] = "active"
    post_title: str | None = Field(default=None, max_length=500)
    keyphrase: str | None = Field(default=None, max_length=500)
    slug: str | None = Field(default=None, max_length=300)
    post_category: str | None = Field(default=None, max_length=200)
    post_url: str | None = Field(default=None, max_length=2000)
    post_status: Literal[
        "empty",
        "defined",
        "creating",
        "published",
        "failed",
    ] = "empty"
    cities: list[PostGroupCityPayload] = Field(default_factory=list)

    def to_repository_payload(self) -> dict[str, Any]:
        payload = self.model_dump(mode="python")
        payload["state"] = self.state.strip().upper()
        payload["cities"] = _normalized_group_cities(self.cities, payload["state"])
        if payload.get("slug"):
            payload["slug"] = str(payload["slug"]).strip("/")
        return payload


class PostGroupUpdateRequest(BaseModel):
    campaign_id: int | None = Field(default=None, gt=0)
    state: str | None = Field(default=None, min_length=2, max_length=120)
    region_name: str | None = Field(default=None, min_length=1, max_length=300)
    group_type: (
        Literal["state_general", "county", "metro", "regional", "custom"] | None
    ) = None
    priority: int | None = Field(default=None, ge=0, le=1000)
    status: Literal["active", "paused", "archived"] | None = None
    post_title: str | None = Field(default=None, max_length=500)
    keyphrase: str | None = Field(default=None, max_length=500)
    slug: str | None = Field(default=None, max_length=300)
    post_category: str | None = Field(default=None, max_length=200)
    post_url: str | None = Field(default=None, max_length=2000)
    post_status: (
        Literal["empty", "defined", "creating", "published", "failed"] | None
    ) = None
    cities: list[PostGroupCityPayload] | None = None

    def to_repository_payload(self) -> dict[str, Any]:
        payload = self.model_dump(mode="python", exclude_unset=True)
        if not payload:
            raise ValueError("At least one field is required")
        if payload.get("state"):
            payload["state"] = str(payload["state"]).strip().upper()
        if payload.get("slug"):
            payload["slug"] = str(payload["slug"]).strip("/")
        if payload.get("cities") is not None:
            payload["cities"] = _normalized_group_cities(
                self.cities or [],
                payload.get("state"),
            )
        return payload


class SchedulePagesRequest(BaseModel):
    campaign_page_ids: list[int] = Field(min_length=1)
    scheduled_for: date
    notes: str = Field(default="Programada desde dashboard", max_length=1000)
    priority: int = Field(default=100, ge=1, le=1000)
    start_now: bool = False
    support_posts: SupportPostsPreference = Field(default_factory=SupportPostsPreference)


class RunNowRequest(BaseModel):
    campaign_page_id: int = Field(gt=0)
    notes: str = Field(
        default="Ejecución manual inmediata desde dashboard",
        max_length=1000,
    )
    support_posts: SupportPostsPreference = Field(default_factory=SupportPostsPreference)


class RunDailyRequest(BaseModel):
    confirm: bool
    daily_limit: int = Field(default=10, ge=1, le=100)
    support_posts: SupportPostsPreference = Field(default_factory=SupportPostsPreference)


class RetryQueueRequest(BaseModel):
    support_posts: SupportPostsPreference = Field(default_factory=SupportPostsPreference)


class ImportPagesRequest(BaseModel):
    filename: str = Field(min_length=1, max_length=255)
    xlsx_base64: str = Field(min_length=1)
    dry_run: bool = True


class CampaignAdminPayload(BaseModel):
    name: str | None = Field(default=None, min_length=1, max_length=200)
    slug: str | None = Field(default=None, min_length=1, max_length=200)
    domain: str | None = Field(default=None, min_length=1, max_length=300)
    business_name: str | None = Field(default=None, min_length=1, max_length=300)
    industry: str | None = Field(default=None, min_length=1, max_length=120)
    language: str | None = Field(default="es", min_length=2, max_length=10)
    main_city: str | None = Field(default=None, max_length=180)
    main_state: str | None = Field(default=None, max_length=120)
    country: str | None = Field(default="US", max_length=10)
    phone: str | None = Field(default=None, max_length=80)
    whatsapp: str | None = Field(default=None, max_length=80)
    address: str | None = Field(default=None, max_length=500)
    business_hours_text: str | None = None
    google_maps_url: str | None = Field(default=None, max_length=2000)
    is_24_hours: bool = False
    contact_page_url: str | None = Field(default=None, max_length=2000)
    contact_page_wp_page_id: int | None = Field(default=None, gt=0)
    years_experience: int | None = Field(default=None, ge=0, le=300)
    experience_text: str | None = None
    brand_phrase: str | None = Field(default=None, max_length=500)
    brand_phrase_exact: bool = False
    brand_phrase_allow_city_variant: bool = True
    brand_phrase_max_uses: int | None = Field(default=None, ge=0, le=50)
    city_links_csv_path: str | None = Field(default=None, max_length=1000)
    button_style_json: dict[str, Any] = Field(default_factory=dict)
    layout_id: int | None = Field(default=None, gt=0)
    layout_branch: Literal["cities", "services"] | None = "cities"
    youtube_channel_id: str | None = Field(
        default=None,
        max_length=120,
        pattern=r"^UC[A-Za-z0-9_-]+$",
    )
    default_elementor_form_id: str | None = Field(default=None, max_length=120)
    active: bool = True

    @field_validator("button_style_json", mode="before")
    @classmethod
    def parse_button_style_json(cls, value: Any) -> Any:
        return _parse_json_object(value, "button_style_json")


class CampaignServiceAdminPayload(BaseModel):
    campaign_id: int | None = Field(default=None, gt=0)
    name: str | None = Field(default=None, min_length=1, max_length=200)
    slug: str | None = Field(default=None, min_length=1, max_length=200)
    primary_keyword: str | None = Field(default=None, max_length=500)
    category: str | None = Field(default=None, max_length=160)
    brand_phrase: str | None = Field(default=None, max_length=500)
    customer_segment: Literal[
        "residential",
        "commercial",
        "both",
        "not_applicable",
        "unknown",
    ] = "unknown"
    language: str | None = Field(default="es", min_length=2, max_length=10)
    active: bool = True
    image_policy: Literal["city_required", "general_allowed", "disabled"] = (
        "general_allowed"
    )
    manual_image_url: str | None = Field(default=None, max_length=2000)
    image_prompt_hint: str | None = None


class WordpressSiteAdminPayload(BaseModel):
    campaign_id: int | None = Field(default=None, gt=0)
    wp_base_url: str | None = Field(default=None, min_length=1, max_length=1000)
    wp_api_base_url: str | None = Field(default=None, max_length=1000)
    wp_admin_url: str | None = Field(default=None, max_length=1000)
    auth_type: str = Field(default="application_password", max_length=80)
    username: str | None = Field(default=None, max_length=200)
    credential_ref: str | None = Field(default=None, max_length=200)
    rest_namespace: str = Field(default="wp/v2", max_length=80)
    elementor_enabled: bool = True
    yoast_enabled: bool = True
    custom_elementor_cache_enabled: bool = False
    automation_key_ref: str | None = Field(default=None, max_length=200)
    active: bool = True


class CampaignPromptRuleAdminPayload(BaseModel):
    campaign_id: int | None = Field(default=None, gt=0)
    page_type: Literal["home", "service", "service_city"] | None = None
    service_id: int | None = Field(default=None, gt=0)
    rules_json: dict[str, Any] = Field(default_factory=dict)
    active: bool = True

    @field_validator("rules_json", mode="before")
    @classmethod
    def parse_rules_json(cls, value: Any) -> Any:
        return _parse_json_object(value, "rules_json")


class SeoManagementService:
    def __init__(
        self,
        repository: SeoAgentManagementRepository,
        page_execution_service: PageExecutionService,
        submit_flow: SubmitFlow,
        rank_position_provider: RankPositionProvider | None = None,
        dispatcher: DispatchExecution | None = None,
    ) -> None:
        self._repository = repository
        self._page_execution_service = page_execution_service
        self._submit_flow = submit_flow
        self._rank_position_provider = rank_position_provider
        self._dispatcher = dispatcher

    async def list_campaigns(self) -> list[dict[str, Any]]:
        return await self._repository.list_campaigns()

    async def list_admin_records(self, resource: AdminResource) -> list[dict[str, Any]]:
        return await self._repository.list_admin_records(resource)

    async def create_admin_record(
        self,
        resource: AdminResource,
        payload: dict[str, Any],
    ) -> dict[str, Any]:
        validated = _validate_admin_payload(resource, payload, partial=False)
        record = await self._repository.create_admin_record(resource, validated)
        return {"ok": True, "record": record}

    async def update_admin_record(
        self,
        resource: AdminResource,
        record_id: int,
        payload: dict[str, Any],
    ) -> dict[str, Any]:
        validated = _validate_admin_payload(resource, payload, partial=True)
        record = await self._repository.update_admin_record(
            resource,
            record_id,
            validated,
        )
        return {"ok": True, "record": record}

    async def delete_admin_record(
        self,
        resource: AdminResource,
        record_id: int,
    ) -> dict[str, Any]:
        record = await self._repository.delete_admin_record(resource, record_id)
        return {"ok": True, "record": record}

    async def list_services(
        self,
        campaign_id: int | None,
        campaign_slug: str | None,
    ) -> list[dict[str, Any]]:
        return await self._repository.list_services(campaign_id, campaign_slug)

    async def list_pages(
        self,
        *,
        campaign_id: int | None,
        campaign_slug: str | None,
        page_type: str | None,
        limit: int,
        offset: int,
    ) -> list[dict[str, Any]]:
        return await self._repository.list_pages(
            campaign_id=campaign_id,
            campaign_slug=campaign_slug,
            page_type=page_type,
            limit=limit,
            offset=offset,
        )

    async def create_page(self, request: CampaignPageCreateRequest) -> dict[str, Any]:
        page = await self._repository.create_page(request.to_repository_payload())
        return {"ok": True, "page": page}

    async def update_page(
        self,
        campaign_page_id: int,
        request: CampaignPageUpdateRequest,
    ) -> dict[str, Any]:
        page = await self._repository.update_page(
            campaign_page_id,
            request.to_repository_payload(),
        )
        return {"ok": True, "page": page}

    async def list_post_groups(
        self,
        *,
        campaign_id: int | None,
        state: str | None,
        post_status: str | None,
        search: str | None,
        limit: int,
    ) -> list[dict[str, Any]]:
        return await self._repository.list_post_groups(
            campaign_id=campaign_id,
            state=state.strip().upper() if state else None,
            post_status=post_status,
            search=search.strip() if search else None,
            limit=limit,
        )

    async def create_post_group(self, request: PostGroupCreateRequest) -> dict[str, Any]:
        group = await self._repository.create_post_group(
            request.to_repository_payload()
        )
        return {"ok": True, "post_group": group}

    async def update_post_group(
        self,
        group_id: int,
        request: PostGroupUpdateRequest,
    ) -> dict[str, Any]:
        group = await self._repository.update_post_group(
            group_id,
            request.to_repository_payload(),
        )
        return {"ok": True, "post_group": group}

    async def delete_post_group(self, group_id: int) -> dict[str, Any]:
        group = await self._repository.delete_post_group(group_id)
        return {"ok": True, "post_group": group}

    async def check_rank_position(self, campaign_page_id: int) -> dict[str, Any]:
        if self._rank_position_provider is None:
            raise RuntimeError("The ranking provider is not configured")

        page = await self._repository.get_page(campaign_page_id)
        if page is None:
            raise LookupError(f"Campaign page {campaign_page_id} was not found")

        keyword = str(page.get("primary_keyword") or "").strip()
        url = str(page.get("url") or "").strip()
        location = str(page.get("target_location_name") or "").strip()
        if not location:
            location = ", ".join(
                str(page.get(field)).strip()
                for field in ("city", "parent_city", "state")
                if page.get(field)
            )
        if not keyword or not url or not location:
            raise ValueError("The page requires a URL, primary keyword, and target location")

        result = await self._rank_position_provider.check_position(
            url=url,
            keyword=keyword,
            location=location,
            country=str(page.get("country") or "US"),
            language=str(page.get("language") or "en"),
        )
        return {
            "ok": True,
            "campaign_page_id": campaign_page_id,
            "keyword": keyword,
            "location": location,
            "engine": "google",
            "result_type": "organic",
            "checked_at": datetime.now(UTC).isoformat(),
            **result,
        }

    async def schedule_pages(self, request: SchedulePagesRequest) -> dict[str, Any]:
        result = await self._repository.schedule_pages(
            request.campaign_page_ids,
            request.scheduled_for,
            _notes_with_support_posts_mode(request.notes, request.support_posts.mode),
            request.priority,
        )
        dispatch_result = {
            "started_count": 0,
            "error_count": 0,
            "started": [],
            "errors": [],
        }
        return {"ok": True, **result, **dispatch_result}

    async def list_queue(
        self,
        *,
        view: str,
        campaign_id: int | None,
        campaign_page_id: int | None,
        page_type: str | None,
        limit: int,
    ) -> list[dict[str, Any]]:
        return await self._repository.list_queue(
            view=view,
            campaign_id=campaign_id,
            campaign_page_id=campaign_page_id,
            page_type=page_type,
            limit=limit,
        )

    async def run_now(self, request: RunNowRequest) -> dict[str, Any]:
        result = await self._repository.create_run_now(
            request.campaign_page_id,
            request.notes,
        )
        if not result.get("ok"):
            return result
        queue_id = int(result["queue"]["id"])
        dispatch = await self._start_queue_flow(
            queue_id,
            request.campaign_page_id,
            request.support_posts.mode,
        )
        return {**result, **dispatch}

    async def retry(
        self,
        queue_execution_id: int,
        request: RetryQueueRequest | None = None,
    ) -> dict[str, Any]:
        result = await self._repository.retry_queue(queue_execution_id)
        queue = result["queue"]
        dispatch = await self._start_queue_flow(
            int(queue["id"]),
            int(queue["campaign_page_id"]),
            (request or RetryQueueRequest()).support_posts.mode,
        )
        return {**result, **dispatch}

    async def cancel(self, queue_execution_id: int) -> dict[str, Any]:
        queue = await self._repository.cancel_queue(queue_execution_id)
        if queue is None:
            raise ValueError("Only queued executions can be cancelled")
        return {"ok": True, "queue": queue}

    async def run_daily(self, request: RunDailyRequest) -> dict[str, Any]:
        if not request.confirm:
            raise ValueError("confirm=true is required")
        queue_items = await self._repository.get_due_queue(request.daily_limit)
        result = await self._start_queue_batch(queue_items, request.support_posts.mode)
        return {"ok": not result["errors"], **result}

    async def _start_queue_batch(
        self,
        queue_items: list[dict[str, Any]],
        default_support_posts_mode: str,
    ) -> dict[str, Any]:
        campaign_bot_ids = await self._assign_campaign_bots(queue_items)
        started: list[dict[str, Any]] = []
        errors: list[dict[str, Any]] = []
        for queue in queue_items:
            queue_id = int(queue["id"])
            page_id = int(queue["campaign_page_id"])
            try:
                support_posts_mode = _support_posts_mode_from_notes(
                    str(queue.get("notes") or ""),
                    default_support_posts_mode,
                )
                dispatch = await self._start_queue_flow(
                    queue_id,
                    page_id,
                    support_posts_mode,
                    preferred_bot_id=campaign_bot_ids.get(_campaign_id_from_queue(queue)),
                )
                started.append({"queue_execution_id": queue_id, **dispatch})
            except Exception as exc:
                errors.append(
                    {
                        "queue_execution_id": queue_id,
                        "error": str(exc),
                    }
                )
        return {
            "started_count": len(started),
            "dispatched_count": sum(
                item["dispatch_status"] in {"queued", "running"} for item in started
            ),
            "pending_bot_count": sum(
                item["dispatch_status"] == "pending" for item in started
            ),
            "already_started_count": sum(
                item["dispatch_status"] == "already_started" for item in started
            ),
            "error_count": len(errors),
            "started": started,
            "errors": errors,
        }

    async def _assign_campaign_bots(
        self,
        queue_items: list[dict[str, Any]],
    ) -> dict[int, UUID]:
        if self._dispatcher is None or not queue_items:
            return {}

        bots = await self._dispatcher.list_online_bots_by_capability("seo.main")
        if not bots:
            return {}

        assignments: dict[int, UUID] = {}
        for queue in queue_items:
            campaign_id = _campaign_id_from_queue(queue)
            if campaign_id is None or campaign_id in assignments:
                continue
            assignments[campaign_id] = bots[len(assignments) % len(bots)].id
        return assignments

    async def import_pages(self, request: ImportPagesRequest) -> dict[str, Any]:
        try:
            raw = base64.b64decode(request.xlsx_base64, validate=False)
        except Exception as exc:
            raise ValueError(f"Invalid XLSX base64: {exc}") from exc
        if len(raw) > 12 * 1024 * 1024:
            raise ValueError("The XLSX file exceeds the 12MB limit")
        try:
            from openpyxl import load_workbook  # type: ignore[import-untyped]
        except ImportError as exc:
            raise RuntimeError("openpyxl is required for XLSX imports") from exc

        try:
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as exc:
            raise ValueError(f"The XLSX file could not be read: {exc}") from exc

        worksheet = (
            workbook["programacion"]
            if "programacion" in workbook.sheetnames
            else workbook.active
        )
        rows = worksheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as exc:
            raise ValueError("The XLSX file has no headers") from exc
        headers = [_normalize_header(value) for value in raw_headers]
        missing = sorted(set(REQUIRED_COLUMNS).difference(headers))
        if missing:
            raise ValueError(f"Missing required columns: {', '.join(missing)}")

        campaigns = await self._repository.list_campaigns()
        campaigns_by_slug = {
            str(campaign.get("slug") or "").strip(): campaign
            for campaign in campaigns
        }
        created: list[dict[str, Any]] = []
        updated: list[dict[str, Any]] = []
        scheduled: list[dict[str, Any]] = []
        skipped: list[dict[str, Any]] = []
        errors: list[dict[str, Any]] = []
        total_rows = 0
        slugs_by_campaign: dict[int, dict[str, int]] = {}

        async def parent_page_id(campaign_id: int, parent_slug: str) -> int | None:
            if campaign_id not in slugs_by_campaign:
                registered = await self._repository.list_pages(
                    campaign_id=campaign_id,
                    limit=1000,
                )
                slugs_by_campaign[campaign_id] = {
                    str(item.get("slug") or "").strip("/").lower(): int(item["id"])
                    for item in registered
                    if item.get("slug") and item.get("id")
                }
            # An unregistered parent is not an error: the slug is stored and the
            # WordPress bot resolves the real parent on the site.
            return slugs_by_campaign[campaign_id].get(parent_slug.strip("/").lower())

        for row_number, raw_row in enumerate(rows, start=2):
            row = {
                header: raw_row[index] if index < len(raw_row) else None
                for index, header in enumerate(headers)
                if header
            }
            if not any(value not in (None, "") for value in row.values()):
                continue
            total_rows += 1
            try:
                campaign_slug = _string(row.get("campaign_slug"))
                campaign = campaigns_by_slug.get(campaign_slug)
                if campaign is None:
                    raise ValueError(f"Unknown campaign_slug: {campaign_slug}")
                campaign_id = int(campaign["id"])
                service_id = None
                service_key = _string(row.get("service_slug"))
                if service_key:
                    services = await self._repository.list_services(
                        campaign_id=campaign_id
                    )
                    service = next(
                        (
                            item
                            for item in services
                            if service_key
                            in {
                                _string(item.get("slug")),
                                _string(item.get("name")),
                            }
                        ),
                        None,
                    )
                    if service is None:
                        raise ValueError(f"Unknown service: {service_key}")
                    service_id = int(service["id"])

                scheduled_for = _date_string(row.get("scheduled_for"))
                if not scheduled_for:
                    raise ValueError("scheduled_for is invalid")
                url = _string(row.get("url"))
                parent_id = _optional_int(row.get("parent_campaign_page_id"))
                parent_slug = _string(row.get("parent_slug")) or None
                if parent_id is None and parent_slug:
                    parent_id = await parent_page_id(campaign_id, parent_slug)
                page_request = CampaignPageCreateRequest(
                    campaign_id=campaign_id,
                    service_id=service_id,
                    wp_page_id=_optional_int(row.get("wp_page_id")),
                    url=url,
                    slug=_string(row.get("slug")) or None,
                    page_type=_string(row.get("page_type")) or "service_city",
                    language=_string(row.get("language")) or "es",
                    primary_keyword=_string(row.get("primary_keyword")),
                    secondary_keywords=_keyword_list(row.get("secondary_keywords")),
                    service_scope=_string(row.get("service_scope")) or "all",
                    customer_segment=_string(row.get("customer_segment")) or "unknown",
                    city=_string(row.get("city")) or None,
                    state=_string(row.get("state")) or None,
                    country=_string(row.get("country")) or "US",
                    target_location_name=_string(row.get("target_location_name")) or None,
                    target_google_maps_url=_string(
                        row.get("target_google_maps_url")
                    )
                    or None,
                    location_type=_string(row.get("location_type")) or "city",
                    parent_city=_string(row.get("parent_city")) or None,
                    delivery_mode=_string(row.get("delivery_mode")) or "unknown",
                    allows_remote=_optional_bool(row.get("allows_remote")),
                    allows_onsite=_optional_bool(row.get("allows_onsite")),
                    physical_visit_required=_optional_bool(
                        row.get("physical_visit_required")
                    ),
                    postal_code=_string(row.get("postal_code")) or None,
                    page_template=_string(row.get("page_template"))
                    or "elementor_header_footer",
                    layout_id=_optional_int(row.get("layout_id")),
                    layout_branch=_string(row.get("layout_branch")) or None,
                    parent_campaign_page_id=parent_id,
                    parent_slug=parent_slug,
                    parent_required=bool(
                        _optional_bool(row.get("parent_required"), False)
                    ),
                    elementor_form_id=_string(row.get("elementor_form_id")) or None,
                    youtube_channel_id=_string(row.get("youtube_channel_id")) or None,
                    allow_publish=_optional_bool(row.get("allow_publish"), True),
                    status=_string(row.get("page_status")) or "active",
                )
                page, action = await self._repository.upsert_import_page(
                    page_request.to_repository_payload(),
                    _optional_int(row.get("campaign_page_id")),
                    request.dry_run,
                )
                item = {
                    "row": row_number,
                    "campaign_page_id": page.get("id"),
                    "url": url,
                    "action": action,
                }
                if "create" in action:
                    created.append(item)
                else:
                    updated.append(item)
                if not request.dry_run and page.get("id") and page.get("slug"):
                    slugs_by_campaign.setdefault(campaign_id, {})[
                        str(page["slug"]).strip("/").lower()
                    ] = int(page["id"])

                if request.dry_run:
                    scheduled.append(
                        {
                            "row": row_number,
                            "campaign_page_id": page.get("id"),
                            "scheduled_for": scheduled_for,
                            "action": "would_schedule",
                        }
                    )
                else:
                    page_id = int(page["id"])
                    schedule_result = await self._repository.schedule_pages(
                        [page_id],
                        date.fromisoformat(scheduled_for),
                        _string(row.get("schedule_notes"))
                        or "Programada desde Excel",
                        _optional_int(row.get("priority")) or 100,
                    )
                    if schedule_result["inserted"]:
                        scheduled.append(
                            {
                                "row": row_number,
                                **schedule_result["inserted"][0],
                            }
                        )
                    else:
                        skipped.append(
                            {
                                "row": row_number,
                                **schedule_result["skipped"][0],
                            }
                        )
            except Exception as exc:
                errors.append({"row": row_number, "error": str(exc)})

        return {
            "ok": True,
            "filename": request.filename,
            "dry_run": request.dry_run,
            "total_rows": total_rows,
            "created_pages_count": len(created),
            "updated_pages_count": len(updated),
            "inserted_queue_count": len(scheduled),
            "skipped_count": len(skipped),
            "error_count": len(errors),
            "created_pages": created[:100],
            "updated_pages": updated[:100],
            "inserted_queue": scheduled[:100],
            "skipped": skipped[:100],
            "errors": errors[:100],
        }

    async def _start_queue_flow(
        self,
        queue_execution_id: int,
        campaign_page_id: int,
        support_posts_mode: str = "auto",
        preferred_bot_id: UUID | None = None,
    ) -> dict[str, Any]:
        page = await self._repository.get_page(campaign_page_id)
        if page is None:
            raise LookupError(f"campaign_page {campaign_page_id} was not found")
        registered_slug = str(page.get("slug") or "").strip().strip("/")
        if not registered_slug:
            raise ValueError(f"campaign_page {campaign_page_id} has no registered slug")
        page_setup_action = _page_setup_action(page)

        await self._page_execution_service.sync_execution(queue_execution_id)
        existing = await self._page_execution_service.get_execution(queue_execution_id)
        existing_flow_id = str((existing or {}).get("orchestrator_flow_id") or "")
        existing_execution_id = str(
            (existing or {}).get("orchestrator_initial_execution_id") or ""
        )
        if existing_flow_id or existing_execution_id:
            if not existing_flow_id or not existing_execution_id:
                raise ValueError(
                    f"Queue execution {queue_execution_id} has an incomplete flow link"
                )
            # A pending bot must not cause another daily run to create a second flow.
            await self._repository.mark_queue_running(queue_execution_id)
            await self._page_execution_service.sync_execution(queue_execution_id)
            return {
                "orchestrator_flow_id": existing_flow_id,
                "orchestrator_execution_id": existing_execution_id,
                "dispatch_status": "already_started",
            }

        request = FlowSubmitRequest(
            flow_type="commercial_page_seo",
            queue_execution_id=queue_execution_id,
            campaign_page_id=campaign_page_id,
            page_url=page["url"],
            campaign=page.get("campaign_name") or page.get("campaign_slug"),
            commercial_objective=page.get("primary_keyword"),
            requested_capability="seo.main",
            correlation_id=f"seo-agent-queue:{queue_execution_id}",
            preferred_bot_id=preferred_bot_id,
            page_setup=WordpressPageSetupRequest(
                action=page_setup_action,
                slug=registered_slug,
                wp_page_id=_optional_int(page.get("wp_page_id")),
                campaign_id=_optional_int(page.get("campaign_id")),
            ),
            payload={
                "campaign_page": page,
                "support_posts": {
                    "mode": support_posts_mode,
                },
            },
        )
        if not await self._repository.mark_queue_running(queue_execution_id):
            raise ValueError(
                f"Queue execution {queue_execution_id} is already starting or no longer queued"
            )
        try:
            response = await self._submit_flow.execute(request)
        except Exception as exc:
            await self._repository.note_queue_start_error(
                queue_execution_id,
                f"Flow startup interrupted: {exc}",
            )
            raise
        dispatch_status = (
            response.dispatched_initial_execution.status
            if response.dispatched_initial_execution is not None
            else response.initial_execution.status
        )
        try:
            await self._page_execution_service.link_orchestrator_flow(
                queue_execution_id,
                str(response.flow.id),
                str(response.initial_execution.id),
            )
        except Exception as exc:
            await self._repository.note_queue_start_error(
                queue_execution_id,
                f"Flow link interrupted: {exc}",
            )
            raise
        if dispatch_status in {ExecutionStatus.FAILED, ExecutionStatus.CANCELLED}:
            dispatched = response.dispatched_initial_execution or response.initial_execution
            await self._repository.update_queue_result(
                queue_execution_id,
                status="failed",
                result_status="failed",
                error_message=dispatched.error_message or "Initial execution could not start",
            )
        await self._page_execution_service.sync_execution(queue_execution_id)
        return {
            "orchestrator_flow_id": str(response.flow.id),
            "orchestrator_execution_id": str(response.initial_execution.id),
            "dispatch_status": dispatch_status.value,
        }


def _slug_from_url(url: str) -> str:
    path = urlparse(url).path.rstrip("/")
    return path.rsplit("/", 1)[-1] if path else "home"


def _city_key(value: str) -> str:
    """Mirror the support_post_city_key() function used by the database.

    support_post_group_cities.city_key is a generated column, so the same
    normalisation has to run here to deduplicate before hitting the unique
    index on (group_id, city_key, state_key).
    """
    text = unicodedata.normalize("NFKD", value.strip().lower())
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = re.sub(r"\s+", " ", text)
    return re.sub(r"[^a-z0-9 ]", "", text).strip()


def _normalized_group_cities(
    cities: list[PostGroupCityPayload],
    default_state: str | None,
) -> list[dict[str, Any]]:
    """Deduplicate the cities of a post group the same way the database does."""
    normalized: dict[tuple[str, str], dict[str, Any]] = {}
    for item in cities:
        city = item.city.strip()
        if not city:
            continue
        state = (item.state or default_state or "").strip().upper()
        if not state:
            raise ValueError(f"City {city} requires a state")
        # The first spelling entered wins over later duplicates.
        normalized.setdefault(
            (_city_key(city), state),
            {"city": city, "state": state},
        )
    return list(normalized.values())


def _page_setup_action(page: dict[str, Any]) -> Literal["create", "update", "skip"]:
    status = str(page.get("status") or "").strip().lower()
    if page.get("has_successful_optimization") is True or status in {
        "optimized",
        "published",
    }:
        return "skip"
    if _optional_int(page.get("wp_page_id")) is not None:
        return "update"
    return "create"


_ADMIN_PAYLOAD_MODELS: dict[str, type[BaseModel]] = {
    "campaigns": CampaignAdminPayload,
    "campaign-services": CampaignServiceAdminPayload,
    "wordpress-sites": WordpressSiteAdminPayload,
    "campaign-prompt-rules": CampaignPromptRuleAdminPayload,
}

_ADMIN_REQUIRED_FIELDS = {
    "campaigns": ("name", "slug", "domain", "business_name", "industry"),
    "campaign-services": ("campaign_id", "name", "slug"),
    "wordpress-sites": ("campaign_id", "wp_base_url"),
    "campaign-prompt-rules": ("campaign_id", "rules_json"),
}


def _validate_admin_payload(
    resource: str,
    payload: dict[str, Any],
    *,
    partial: bool,
) -> dict[str, Any]:
    model_type = _ADMIN_PAYLOAD_MODELS.get(resource)
    if model_type is None:
        raise ValueError(f"Unsupported admin resource: {resource}")

    try:
        model = model_type.model_validate(payload)
    except ValidationError as exc:
        raise ValueError(str(exc)) from exc
    data = model.model_dump(mode="python", exclude_unset=partial)
    data = {
        key: (value.strip() if isinstance(value, str) else value)
        for key, value in data.items()
    }
    data = {key: (None if value == "" else value) for key, value in data.items()}
    if not partial:
        missing = [
            field
            for field in _ADMIN_REQUIRED_FIELDS[resource]
            if data.get(field) in (None, "", {})
        ]
        if missing:
            raise ValueError(f"Missing required fields: {', '.join(missing)}")
    if partial and not data:
        raise ValueError("At least one field is required")
    return data


def _parse_json_object(value: Any, field_name: str) -> dict[str, Any]:
    if value in (None, ""):
        return {}
    if isinstance(value, str):
        try:
            value = json.loads(value)
        except json.JSONDecodeError as exc:
            raise ValueError(f"{field_name} must be valid JSON: {exc}") from exc
    if not isinstance(value, dict):
        raise ValueError(f"{field_name} must be a JSON object")
    return value


def _notes_with_support_posts_mode(notes: str, mode: str) -> str:
    clean_notes = notes.strip()
    marker = f"[support_posts.mode={mode}]"
    return f"{clean_notes} {marker}".strip()


def _support_posts_mode_from_notes(notes: str, default: str = "auto") -> str:
    fallback = default if default in SUPPORT_POST_MODES else "auto"
    marker = "[support_posts.mode="
    start = notes.find(marker)
    if start == -1:
        return fallback
    value_start = start + len(marker)
    value_end = notes.find("]", value_start)
    if value_end == -1:
        return fallback
    mode = notes[value_start:value_end]
    return mode if mode in SUPPORT_POST_MODES else fallback


def _campaign_id_from_queue(queue: dict[str, Any]) -> int | None:
    value = queue.get("campaign_id")
    if value is None:
        return None
    try:
        campaign_id = int(value)
    except (TypeError, ValueError):
        return None
    return campaign_id if campaign_id > 0 else None


_HEADER_ALIASES = {
    "campaign": "campaign_slug",
    "campania": "campaign_slug",
    "campaÃ±a": "campaign_slug",
    "service": "service_slug",
    "service_name": "service_slug",
    "servicio": "service_slug",
    "keyword": "primary_keyword",
    "main_keyword": "primary_keyword",
    "fecha": "scheduled_for",
    "fecha_programada": "scheduled_for",
    "programada_para": "scheduled_for",
    "status": "page_status",
    "notes": "schedule_notes",
    "notas": "schedule_notes",
    "page_id": "wp_page_id",
    "zip": "postal_code",
    "zip_code": "postal_code",
    "codigo_postal": "postal_code",
    "parent": "parent_slug",
    "padre": "parent_slug",
    "pagina_padre": "parent_slug",
    "parent_page_slug": "parent_slug",
    "parent_page_id": "parent_campaign_page_id",
    "plantilla": "page_template",
    "template": "page_template",
    "formulario": "elementor_form_id",
    "form_id": "elementor_form_id",
    "youtube": "youtube_channel_id",
    "canal_youtube": "youtube_channel_id",
    "prioridad": "priority",
}


def _normalize_header(value: Any) -> str:
    header = _string(value).lower().replace("\ufeff", "")
    for character in (" ", "-", ".", "/", "\\"):
        header = header.replace(character, "_")
    while "__" in header:
        header = header.replace("__", "_")
    normalized = header.strip("_")
    return _HEADER_ALIASES.get(normalized, normalized)


def _string(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _optional_int(value: Any) -> int | None:
    string = _string(value)
    return int(float(string)) if string else None


def _optional_bool(value: Any, default: bool | None = None) -> bool | None:
    string = _string(value).lower()
    if not string:
        return default
    if string in {"true", "1", "yes", "y", "si", "sí"}:
        return True
    if string in {"false", "0", "no", "n"}:
        return False
    return default


def _keyword_list(value: Any) -> list[str]:
    if isinstance(value, list):
        return [_string(item) for item in value if _string(item)]
    string = _string(value)
    if not string:
        return []
    return [item.strip() for item in string.replace("\n", ",").split(",") if item.strip()]


def _date_string(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    string = _string(value)
    if not string:
        return ""
    candidate = string.split("T", 1)[0].split(" ", 1)[0]
    for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(candidate, pattern).date().isoformat()
        except ValueError:
            continue
    return ""
