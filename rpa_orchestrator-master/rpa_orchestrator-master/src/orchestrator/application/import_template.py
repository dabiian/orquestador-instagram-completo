"""Single source of truth for the pages/scheduling XLSX import template.

The same column list feeds the downloadable template and the documentation, so
the file the user fills in always matches what ``SeoManagementService.import_pages``
reads.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any

IMPORT_TEMPLATE_FILENAME = "programacion_paginas_template.xlsx"
SHEET_NAME = "programacion"
GUIDE_SHEET_NAME = "guia"
TEMPLATE_ROWS = 200


@dataclass(frozen=True)
class TemplateColumn:
    name: str
    group: str
    required: bool
    description: str
    example: str = ""
    options: tuple[str, ...] = field(default_factory=tuple)

    @property
    def valid_values(self) -> str:
        return " | ".join(self.options) if self.options else ""


IMPORT_TEMPLATE_COLUMNS: tuple[TemplateColumn, ...] = (
    TemplateColumn(
        "campaign_slug",
        "Obligatorias",
        True,
        "Slug de la campaña ya registrada en campaigns.",
        "boxmarkdigital-com",
    ),
    TemplateColumn(
        "url",
        "Obligatorias",
        True,
        "URL completa y definitiva de la página.",
        "https://ejemplo.com/fences/aurora/",
    ),
    TemplateColumn(
        "page_type",
        "Obligatorias",
        True,
        "Tipo de página dentro de la campaña.",
        "service_city",
        ("home", "service", "service_city"),
    ),
    TemplateColumn(
        "primary_keyword",
        "Obligatorias",
        True,
        "Keyword principal que debe posicionar la página.",
        "Fence Company Aurora",
    ),
    TemplateColumn(
        "scheduled_for",
        "Obligatorias",
        True,
        "Fecha de la ejecución programada (AAAA-MM-DD).",
        "2026-09-25",
    ),
    TemplateColumn(
        "campaign_page_id",
        "Identificación",
        False,
        "ID existente en campaign_pages. Si se indica, actualiza esa fila.",
        "1421",
    ),
    TemplateColumn(
        "wp_page_id",
        "Identificación",
        False,
        "ID de la página en WordPress. Vacío = el bot la crea; con valor = la actualiza.",
        "5120",
    ),
    TemplateColumn(
        "slug",
        "Identificación",
        False,
        "Slug registrado. Si se omite se deriva de la URL. Ningún bot lo cambia.",
        "fences-aurora",
    ),
    TemplateColumn(
        "language",
        "Identificación",
        False,
        "Idioma del contenido.",
        "es",
        ("es", "en"),
    ),
    TemplateColumn(
        "service_slug",
        "Identificación",
        False,
        "Slug o nombre del servicio en campaign_services.",
        "fences",
    ),
    TemplateColumn(
        "service_scope",
        "Identificación",
        False,
        "Si la página cubre todos los servicios o solo uno.",
        "all",
        ("all", "single"),
    ),
    TemplateColumn(
        "secondary_keywords",
        "Keywords",
        False,
        "Keywords secundarias separadas por coma o salto de línea.",
        "fence installation aurora, vinyl fence aurora",
    ),
    TemplateColumn(
        "city",
        "Ubicación",
        False,
        "Ciudad o barrio objetivo de la página.",
        "Aurora",
    ),
    TemplateColumn("state", "Ubicación", False, "Estado o región.", "IL"),
    TemplateColumn("country", "Ubicación", False, "País (ISO).", "US"),
    TemplateColumn(
        "postal_code",
        "Ubicación",
        False,
        "Código postal que se inyecta en el schema local de la página.",
        "60505",
    ),
    TemplateColumn(
        "target_location_name",
        "Ubicación",
        False,
        "Nombre de ubicación tal como debe aparecer en el contenido.",
        "Aurora, IL",
    ),
    TemplateColumn(
        "target_google_maps_url",
        "Ubicación",
        False,
        "URL de Google Maps de la ubicación objetivo.",
        "https://maps.app.goo.gl/xxxx",
    ),
    TemplateColumn(
        "location_type",
        "Ubicación",
        False,
        "Tipo de ubicación. 'neighborhood' exige parent_city.",
        "city",
        ("city", "neighborhood", "general"),
    ),
    TemplateColumn(
        "parent_city",
        "Ubicación",
        False,
        "Ciudad principal de un barrio. Relación geográfica, no jerarquía de WordPress.",
        "Chicago",
    ),
    TemplateColumn(
        "delivery_mode",
        "Operación",
        False,
        "Modalidad de prestación del servicio.",
        "onsite",
        ("remote", "onsite", "hybrid", "unknown"),
    ),
    TemplateColumn(
        "customer_segment",
        "Operación",
        False,
        "Segmento de cliente al que apunta la página.",
        "commercial",
        ("residential", "commercial", "both", "not_applicable", "unknown"),
    ),
    TemplateColumn(
        "allows_remote",
        "Operación",
        False,
        "Si el servicio se puede prestar a distancia.",
        "false",
        ("true", "false"),
    ),
    TemplateColumn(
        "allows_onsite",
        "Operación",
        False,
        "Si el servicio se presta en el domicilio del cliente.",
        "true",
        ("true", "false"),
    ),
    TemplateColumn(
        "physical_visit_required",
        "Operación",
        False,
        "Si se requiere visita física obligatoria.",
        "true",
        ("true", "false"),
    ),
    TemplateColumn(
        "allow_publish",
        "Operación",
        False,
        "Si el bot puede publicar automáticamente al terminar.",
        "true",
        ("true", "false"),
    ),
    TemplateColumn(
        "parent_campaign_page_id",
        "Jerarquía WordPress",
        False,
        "ID en campaign_pages de la página padre. Tiene prioridad sobre parent_slug.",
        "1330",
    ),
    TemplateColumn(
        "parent_slug",
        "Jerarquía WordPress",
        False,
        "Slug de la página padre. El sistema resuelve el parent_wp_page_id.",
        "fences",
    ),
    TemplateColumn(
        "parent_required",
        "Jerarquía WordPress",
        False,
        "Si el padre debe existir antes de crear esta página.",
        "false",
        ("true", "false"),
    ),
    TemplateColumn(
        "page_template",
        "Plantilla y formularios",
        False,
        "Plantilla de WordPress asignada a la página.",
        "elementor_header_footer",
    ),
    TemplateColumn(
        "elementor_form_id",
        "Plantilla y formularios",
        False,
        "Solo para sobrescribir el formulario por defecto de la campaña.",
        "a1b2c3d",
    ),
    TemplateColumn(
        "youtube_channel_id",
        "Plantilla y formularios",
        False,
        "Solo para sobrescribir el canal de YouTube de la campaña.",
        "UCxxxxxxxxxxxxxxxxxxxxxx",
    ),
    TemplateColumn(
        "layout_id",
        "Plantilla y formularios",
        False,
        "Solo si la página usa un layout distinto al de la campaña.",
        "23",
    ),
    TemplateColumn(
        "layout_branch",
        "Plantilla y formularios",
        False,
        "Rama del layout cuando se sobrescribe el de la campaña.",
        "cities",
        ("cities", "services"),
    ),
    TemplateColumn(
        "page_status",
        "Programación",
        False,
        "Estado inicial de la página en la base.",
        "active",
        ("pending", "active", "published"),
    ),
    TemplateColumn(
        "priority",
        "Programación",
        False,
        "Prioridad en la cola. Menor número se ejecuta antes.",
        "100",
    ),
    TemplateColumn(
        "schedule_notes",
        "Programación",
        False,
        "Nota que queda registrada en page_execution_queue.",
        "Lote septiembre",
    ),
)

REQUIRED_COLUMNS: tuple[str, ...] = tuple(
    column.name for column in IMPORT_TEMPLATE_COLUMNS if column.required
)

_GUIDE_NOTES: tuple[str, ...] = (
    "El canal de YouTube, el formulario Elementor y el layout se heredan de la "
    "campaña. Complete esas columnas solo para excepciones de una página.",
    "La acción crear / actualizar / omitir no se indica en el Excel: se deduce de "
    "wp_page_id y del historial de la página.",
    "parent_city es una relación geográfica y no reemplaza a parent_slug ni a "
    "parent_campaign_page_id, que definen la jerarquía real de WordPress.",
    "parent_wp_page_id lo resuelve el sistema; no existe como columna del Excel.",
    "Si parent_slug no corresponde a ninguna página registrada, el slug se guarda "
    "igual y el bot de WordPress lo resuelve en el sitio.",
    "No cambie los encabezados de la fila 1 de la hoja 'programacion'.",
)


def build_import_template() -> bytes:
    """Return the XLSX import template as bytes."""
    from openpyxl import Workbook  # type: ignore[import-untyped]
    from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
    from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
    from openpyxl.worksheet.datavalidation import (  # type: ignore[import-untyped]
        DataValidation,
    )

    group_colors = {
        "Obligatorias": "1F4E79",
        "Identificación": "2E7D32",
        "Keywords": "6A1B9A",
        "Ubicación": "B26A00",
        "Operación": "455A64",
        "Jerarquía WordPress": "AD1457",
        "Plantilla y formularios": "00695C",
        "Programación": "37474F",
    }

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.freeze_panes = "A2"

    for index, column in enumerate(IMPORT_TEMPLATE_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=column.name)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill(
            "solid",
            fgColor=group_colors.get(column.group, "37474F"),
        )
        cell.alignment = Alignment(horizontal="center", vertical="center")
        comment_lines = [column.group, column.description]
        if column.valid_values:
            comment_lines.append(f"Valores: {column.valid_values}")
        if column.example:
            comment_lines.append(f"Ejemplo: {column.example}")
        _set_note(cell, "\n".join(comment_lines))
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = max(14, min(len(column.name) + 8, 34))
        if column.options:
            validation = DataValidation(
                type="list",
                formula1=f'"{",".join(column.options)}"',
                allow_blank=True,
                showDropDown=False,
            )
            sheet.add_data_validation(validation)
            validation.add(f"{letter}2:{letter}{TEMPLATE_ROWS + 1}")

    guide = workbook.create_sheet(GUIDE_SHEET_NAME)
    guide.append(["Plantilla de importación de páginas y programación"])
    guide["A1"].font = Font(bold=True, size=14)
    guide.append([])
    guide.append(["Grupo", "Columna", "Obligatoria", "Descripción", "Ejemplo", "Valores válidos"])
    for cell in guide[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
    for column in IMPORT_TEMPLATE_COLUMNS:
        guide.append(
            [
                column.group,
                column.name,
                "Sí" if column.required else "No",
                column.description,
                column.example,
                column.valid_values,
            ]
        )
    guide.append([])
    guide.append(["Notas"])
    guide.cell(row=guide.max_row, column=1).font = Font(bold=True, size=12)
    for note in _GUIDE_NOTES:
        guide.append([note])
    for letter, width in (
        ("A", 24),
        ("B", 26),
        ("C", 12),
        ("D", 72),
        ("E", 34),
        ("F", 52),
    ):
        guide.column_dimensions[letter].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _set_note(cell: Any, text: str) -> None:
    from openpyxl.comments import Comment  # type: ignore[import-untyped]

    comment = Comment(text, "SEO Orchestrator")
    comment.width = 320
    comment.height = 130
    cell.comment = comment
