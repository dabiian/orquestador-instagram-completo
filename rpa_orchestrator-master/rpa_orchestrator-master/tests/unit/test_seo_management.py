from __future__ import annotations

import base64
import io
from datetime import date
from types import SimpleNamespace
from typing import Any
from uuid import uuid4

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]

from orchestrator.api.v1.dependencies import seo_management_service
from orchestrator.api.v1.routes.seo_management import router
from orchestrator.application.import_template import (
    IMPORT_TEMPLATE_FILENAME,
    REQUIRED_COLUMNS,
    SHEET_NAME,
)
from orchestrator.application.seo_management import (
    CampaignPageCreateRequest,
    CampaignPageUpdateRequest,
    ImportPagesRequest,
    PostGroupCreateRequest,
    PostGroupUpdateRequest,
    RunDailyRequest,
    SchedulePagesRequest,
    SeoManagementService,
    _page_setup_action,
)
from orchestrator.domain.entities import Bot, ExecutionStatus


class FakeManagementRepository:
    def __init__(self) -> None:
        self.page_filters: dict[str, Any] = {}
        self.created_payload: dict[str, Any] | None = None
        self.scheduled: list[int] = []
        self.scheduled_for: date | None = None
        self.schedule_notes: str | None = None
        self.schedule_priority: int | None = None
        self.page: dict[str, Any] | None = None
        self.pages: dict[int, dict[str, Any]] = {}
        self.updated_page_payload: dict[str, Any] | None = None
        self.post_groups: list[dict[str, Any]] = []
        self.post_group_filters: dict[str, Any] = {}
        self.post_group_payload: dict[str, Any] | None = None
        self.admin_records: dict[str, list[dict[str, Any]]] = {
            "campaigns": [{"id": 1, "name": "Campaign One", "slug": "campaign-one"}],
            "campaign-services": [],
            "wordpress-sites": [],
            "campaign-prompt-rules": [],
        }
        self.admin_payload: dict[str, Any] | None = None
        self.running_queue_ids: list[int] = []
        self.claimed_queue_ids: set[int] = set()
        self.queue_results: list[dict[str, Any]] = []
        self.start_errors: list[tuple[int, str]] = []

    async def list_campaigns(self) -> list[dict[str, Any]]:
        return [{"id": 1, "slug": "campaign-one", "name": "Campaign One"}]

    async def list_admin_records(self, resource: str) -> list[dict[str, Any]]:
        return self.admin_records[resource]

    async def create_admin_record(
        self,
        resource: str,
        payload: dict[str, Any],
    ) -> dict[str, Any]:
        self.admin_payload = payload
        record = {"id": 99, **payload}
        self.admin_records.setdefault(resource, []).append(record)
        return record

    async def update_admin_record(
        self,
        resource: str,
        record_id: int,
        payload: dict[str, Any],
    ) -> dict[str, Any]:
        self.admin_payload = payload
        return {"id": record_id, **payload}

    async def delete_admin_record(
        self,
        resource: str,
        record_id: int,
    ) -> dict[str, Any]:
        return {"id": record_id, "deleted_from": resource}

    async def list_services(
        self,
        campaign_id: int | None = None,
        campaign_slug: str | None = None,
    ) -> list[dict[str, Any]]:
        return []

    async def list_pages(self, **kwargs: Any) -> list[dict[str, Any]]:
        self.page_filters = kwargs
        return [{"id": 188, "campaign_id": kwargs.get("campaign_id")}]

    async def create_page(self, payload: dict[str, Any]) -> dict[str, Any]:
        self.created_payload = payload
        return {"id": 188, **payload}

    async def update_page(
        self,
        campaign_page_id: int,
        payload: dict[str, Any],
    ) -> dict[str, Any]:
        self.updated_page_payload = payload
        page = self.pages.get(campaign_page_id, {"id": campaign_page_id})
        page.update(payload)
        return page

    async def list_post_groups(
        self,
        *,
        campaign_id: int | None = None,
        state: str | None = None,
        post_status: str | None = None,
        search: str | None = None,
        limit: int = 500,
    ) -> list[dict[str, Any]]:
        self.post_group_filters = {
            "campaign_id": campaign_id,
            "state": state,
            "post_status": post_status,
            "search": search,
            "limit": limit,
        }
        return self.post_groups

    async def create_post_group(self, payload: dict[str, Any]) -> dict[str, Any]:
        self.post_group_payload = payload
        group = {"id": 501, **payload}
        self.post_groups.append(group)
        return group

    async def update_post_group(
        self,
        group_id: int,
        payload: dict[str, Any],
    ) -> dict[str, Any]:
        self.post_group_payload = payload
        return {"id": group_id, **payload}

    async def delete_post_group(self, group_id: int) -> dict[str, Any]:
        return {"id": group_id}

    async def upsert_import_page(
        self,
        payload: dict[str, Any],
        campaign_page_id: int | None,
        dry_run: bool,
    ) -> tuple[dict[str, Any], str]:
        self.created_payload = payload
        return {"id": None if dry_run else 188, **payload}, (
            "would_create" if dry_run else "created"
        )

    async def schedule_pages(
        self,
        campaign_page_ids: list[int],
        scheduled_for: date,
        notes: str,
        priority: int = 100,
    ) -> dict[str, Any]:
        self.scheduled.extend(campaign_page_ids)
        self.scheduled_for = scheduled_for
        self.schedule_notes = notes
        self.schedule_priority = priority
        return {
            "inserted_count": len(campaign_page_ids),
            "inserted": [
                {
                    "id": 36 + index,
                    "campaign_page_id": campaign_page_id,
                    "campaign_id": self.pages.get(campaign_page_id, {}).get("campaign_id", 1),
                    "scheduled_for": scheduled_for,
                }
                for index, campaign_page_id in enumerate(campaign_page_ids)
            ],
            "skipped": [],
        }

    async def list_queue(self, **kwargs: Any) -> list[dict[str, Any]]:
        return []

    async def create_run_now(
        self,
        campaign_page_id: int,
        notes: str,
    ) -> dict[str, Any]:
        raise NotImplementedError

    async def retry_queue(self, queue_execution_id: int) -> dict[str, Any]:
        raise NotImplementedError

    async def cancel_queue(self, queue_execution_id: int) -> dict[str, Any] | None:
        raise NotImplementedError

    async def get_page(self, campaign_page_id: int) -> dict[str, Any] | None:
        if campaign_page_id in self.pages:
            return self.pages[campaign_page_id]
        return self.page if self.page and self.page["id"] == campaign_page_id else None

    async def update_page_wp_page_id(
        self,
        campaign_page_id: int,
        wp_page_id: int,
    ) -> None:
        page = await self.get_page(campaign_page_id)
        if page is not None:
            page["wp_page_id"] = wp_page_id

    async def get_due_queue(self, limit: int) -> list[dict[str, Any]]:
        return []

    async def mark_queue_running(self, queue_execution_id: int) -> bool:
        self.running_queue_ids.append(queue_execution_id)
        if queue_execution_id in self.claimed_queue_ids:
            return False
        self.claimed_queue_ids.add(queue_execution_id)
        return True

    async def note_queue_start_error(
        self, queue_execution_id: int, error_message: str
    ) -> None:
        self.start_errors.append((queue_execution_id, error_message))

    async def update_queue_result(
        self,
        queue_execution_id: int,
        *,
        status: str,
        result_status: str,
        error_message: str | None = None,
    ) -> None:
        self.queue_results.append(
            {
                "queue_execution_id": queue_execution_id,
                "status": status,
                "result_status": result_status,
                "error_message": error_message,
            }
        )


class FakeRankPositionProvider:
    def __init__(self) -> None:
        self.request: dict[str, str] | None = None

    async def check_position(self, **kwargs: str) -> dict[str, Any]:
        self.request = kwargs
        return {
            "request_id": "rank-request-1",
            "position": 7,
            "matched_url": "https://example.com/my-page/",
            "max_results": 100,
        }


class FakePageExecutionService:
    def __init__(self) -> None:
        self.synced: list[int] = []
        self.linked: list[tuple[int, str, str]] = []
        self.executions: dict[int, dict[str, Any]] = {}

    async def sync_execution(self, queue_execution_id: int) -> None:
        self.synced.append(queue_execution_id)

    async def get_execution(self, queue_execution_id: int) -> dict[str, Any] | None:
        return self.executions.get(queue_execution_id)

    async def link_orchestrator_flow(
        self,
        queue_execution_id: int,
        flow_id: str,
        execution_id: str,
    ) -> None:
        self.linked.append((queue_execution_id, flow_id, execution_id))
        self.executions[queue_execution_id] = {
            "orchestrator_flow_id": flow_id,
            "orchestrator_initial_execution_id": execution_id,
        }


class FakeSubmitFlow:
    def __init__(self) -> None:
        self.requests: list[Any] = []

    async def execute(self, request: Any) -> Any:
        self.requests.append(request)
        return SimpleNamespace(
            flow=SimpleNamespace(id=uuid4()),
            initial_execution=SimpleNamespace(id=uuid4()),
            dispatched_initial_execution=SimpleNamespace(status=ExecutionStatus.QUEUED),
        )


class FakeDispatcher:
    def __init__(self, bots: list[Bot]) -> None:
        self.bots = bots

    async def list_online_bots_by_capability(self, capability: str) -> list[Bot]:
        return [bot for bot in self.bots if capability in bot.capabilities]


def make_service(
    repository: FakeManagementRepository,
    rank_position_provider: FakeRankPositionProvider | None = None,
    page_execution_service: Any = None,
    submit_flow: Any = None,
    dispatcher: Any = None,
) -> SeoManagementService:
    return SeoManagementService(
        repository=repository,
        page_execution_service=page_execution_service or object(),  # type: ignore[arg-type]
        submit_flow=submit_flow or object(),  # type: ignore[arg-type]
        rank_position_provider=rank_position_provider,
        dispatcher=dispatcher,
    )


async def test_create_page_derives_slug_and_keeps_campaign() -> None:
    repository = FakeManagementRepository()
    service = make_service(repository)

    result = await service.create_page(
        CampaignPageCreateRequest(
            campaign_id=1,
            url="https://example.com/my-page/",
            primary_keyword="My keyword",
        )
    )

    assert result["page"]["slug"] == "my-page"
    assert repository.created_payload is not None
    assert repository.created_payload["campaign_id"] == 1


def test_page_setup_action_uses_previous_optimization_and_wp_identity() -> None:
    assert _page_setup_action({"has_successful_optimization": True}) == "skip"
    assert _page_setup_action({"status": "published", "wp_page_id": 10}) == "skip"
    assert _page_setup_action({"status": "active", "wp_page_id": 10}) == "update"
    assert _page_setup_action({"status": "pending", "wp_page_id": None}) == "create"


def test_registered_pages_endpoint_filters_by_campaign() -> None:
    repository = FakeManagementRepository()
    service = make_service(repository)
    app = FastAPI()
    app.include_router(router, prefix="/seo")
    app.dependency_overrides[seo_management_service] = lambda: service

    with TestClient(app) as client:
        response = client.get("/seo/pages?campaign_id=7")

    assert response.status_code == 200
    assert repository.page_filters["campaign_id"] == 7


async def test_create_admin_campaign_service_validates_and_persists_payload() -> None:
    repository = FakeManagementRepository()
    service = make_service(repository)

    result = await service.create_admin_record(
        "campaign-services",
        {
            "campaign_id": 1,
            "name": "Office Cleaning",
            "slug": "office-cleaning",
            "customer_segment": "commercial",
            "image_policy": "general_allowed",
        },
    )

    assert result["record"]["id"] == 99
    assert repository.admin_payload is not None
    assert repository.admin_payload["campaign_id"] == 1
    assert repository.admin_payload["active"] is True


async def test_admin_prompt_rules_requires_json_object() -> None:
    repository = FakeManagementRepository()
    service = make_service(repository)

    try:
        await service.create_admin_record(
            "campaign-prompt-rules",
            {"campaign_id": 1, "rules_json": ["not", "an", "object"]},
        )
    except ValueError as exc:
        assert "rules_json must be a JSON object" in str(exc)
    else:
        raise AssertionError("Expected ValueError")


def test_admin_delete_endpoint_returns_json() -> None:
    repository = FakeManagementRepository()
    service = make_service(repository)
    app = FastAPI()
    app.include_router(router, prefix="/seo")
    app.dependency_overrides[seo_management_service] = lambda: service

    with TestClient(app) as client:
        response = client.delete("/seo/admin/campaigns/1")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/json")
    assert response.json()["record"] == {"id": 1, "deleted_from": "campaigns"}


def test_create_page_endpoint_serializes_repository_dates() -> None:
    class DateRepository(FakeManagementRepository):
        async def create_page(self, payload: dict[str, Any]) -> dict[str, Any]:
            self.created_payload = payload
            return {"id": 188, "created_at": date(2026, 7, 16), **payload}

    repository = DateRepository()
    service = make_service(repository)
    app = FastAPI()
    app.include_router(router, prefix="/seo")
    app.dependency_overrides[seo_management_service] = lambda: service

    with TestClient(app) as client:
        response = client.post(
            "/seo/pages",
            json={
                "campaign_id": 1,
                "url": "https://example.com/my-page/",
                "primary_keyword": "My keyword",
            },
        )

    assert response.status_code == 201
    assert response.headers["content-type"].startswith("application/json")
    assert response.json()["page"]["created_at"] == "2026-07-16"


def test_seo_endpoint_unexpected_errors_are_json() -> None:
    class ErrorRepository(FakeManagementRepository):
        async def create_page(self, payload: dict[str, Any]) -> dict[str, Any]:
            raise TypeError("could not serialize database row")

    repository = ErrorRepository()
    service = make_service(repository)
    app = FastAPI()
    app.include_router(router, prefix="/seo")
    app.dependency_overrides[seo_management_service] = lambda: service

    with TestClient(app) as client:
        response = client.post(
            "/seo/pages",
            json={
                "campaign_id": 1,
                "url": "https://example.com/my-page/",
                "primary_keyword": "My keyword",
            },
        )

    assert response.status_code == 500
    assert response.headers["content-type"].startswith("application/json")
    assert response.json() == {"detail": "could not serialize database row"}


def test_schedule_endpoint_serializes_inserted_dates() -> None:
    class DateScheduleRepository(FakeManagementRepository):
        async def schedule_pages(
            self,
            campaign_page_ids: list[int],
            scheduled_for: str,
            notes: str,
            priority: int = 100,
        ) -> dict[str, Any]:
            self.scheduled.extend(campaign_page_ids)
            self.schedule_notes = notes
            self.schedule_priority = priority
            return {
                "inserted_count": 1,
                "inserted": [
                    {
                        "id": 36,
                        "campaign_page_id": campaign_page_ids[0],
                        "campaign_id": 1,
                        "scheduled_for": date(2026, 7, 16),
                    }
                ],
                "skipped": [],
            }

    repository = DateScheduleRepository()
    service = make_service(repository)
    app = FastAPI()
    app.include_router(router, prefix="/seo")
    app.dependency_overrides[seo_management_service] = lambda: service

    with TestClient(app) as client:
        response = client.post(
            "/seo/execution-queue/schedule",
            json={
                "campaign_page_ids": [188],
                "scheduled_for": "2026-07-16",
                "start_now": False,
            },
        )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/json")
    assert response.json()["inserted"][0]["scheduled_for"] == "2026-07-16"


async def test_check_rank_position_uses_registered_page_data() -> None:
    repository = FakeManagementRepository()
    repository.page = {
        "id": 188,
        "url": "https://example.com/my-page/",
        "primary_keyword": "Local service",
        "target_location_name": "Bogotá, Colombia",
        "country": "CO",
        "language": "es",
    }
    provider = FakeRankPositionProvider()
    service = make_service(repository, provider)

    result = await service.check_rank_position(188)

    assert result["position"] == 7
    assert result["result_type"] == "organic"
    assert provider.request == {
        "url": "https://example.com/my-page/",
        "keyword": "Local service",
        "location": "Bogotá, Colombia",
        "country": "CO",
        "language": "es",
    }


async def test_xlsx_import_creates_and_schedules_page() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(
        [
            "campaign_slug",
            "url",
            "page_type",
            "primary_keyword",
            "scheduled_for",
        ]
    )
    worksheet.append(
        [
            "campaign-one",
            "https://example.com/local-service/",
            "service_city",
            "Local service",
            "2026-07-05",
        ]
    )
    buffer = io.BytesIO()
    workbook.save(buffer)

    repository = FakeManagementRepository()
    service = make_service(repository)
    result = await service.import_pages(
        ImportPagesRequest(
            filename="pages.xlsx",
            xlsx_base64=base64.b64encode(buffer.getvalue()).decode("ascii"),
            dry_run=False,
        )
    )

    assert result["created_pages_count"] == 1
    assert result["inserted_queue_count"] == 1
    assert result["error_count"] == 0
    assert repository.scheduled == [188]


async def test_xlsx_import_maps_wordpress_hierarchy_and_priority() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(
        [
            "campaign_slug",
            "url",
            "page_type",
            "primary_keyword",
            "scheduled_for",
            "postal_code",
            "parent_slug",
            "parent_required",
            "page_template",
            "elementor_form_id",
            "youtube_channel_id",
            "layout_id",
            "layout_branch",
            "priority",
        ]
    )
    worksheet.append(
        [
            "campaign-one",
            "https://example.com/fences/aurora/",
            "service_city",
            "Fence Company Aurora",
            "2026-09-25",
            "60505",
            "fences",
            "true",
            "elementor_full_width",
            "a1b2c3d",
            "UCtest",
            "23",
            "cities",
            "5",
        ]
    )
    buffer = io.BytesIO()
    workbook.save(buffer)

    repository = FakeManagementRepository()
    service = make_service(repository)
    result = await service.import_pages(
        ImportPagesRequest(
            filename="pages.xlsx",
            xlsx_base64=base64.b64encode(buffer.getvalue()).decode("ascii"),
            dry_run=False,
        )
    )

    assert result["error_count"] == 0
    payload = repository.created_payload
    assert payload is not None
    assert payload["postal_code"] == "60505"
    assert payload["parent_slug"] == "fences"
    assert payload["parent_required"] is True
    assert payload["page_template"] == "elementor_full_width"
    assert payload["elementor_form_id"] == "a1b2c3d"
    assert payload["youtube_channel_id"] == "UCtest"
    assert payload["layout_id"] == 23
    assert payload["layout_branch"] == "cities"
    assert "parent_wp_page_id" not in payload
    assert repository.schedule_priority == 5


async def test_xlsx_import_resolves_parent_slug_to_registered_page() -> None:
    class ParentAwareRepository(FakeManagementRepository):
        async def list_pages(self, **kwargs: Any) -> list[dict[str, Any]]:
            self.page_filters = kwargs
            return [{"id": 700, "slug": "Fences"}]

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(
        [
            "campaign_slug",
            "url",
            "page_type",
            "primary_keyword",
            "scheduled_for",
            "parent_slug",
        ]
    )
    worksheet.append(
        [
            "campaign-one",
            "https://example.com/fences/aurora/",
            "service_city",
            "Fence Company Aurora",
            "2026-09-25",
            "fences",
        ]
    )
    worksheet.append(
        [
            "campaign-one",
            "https://example.com/fences/naperville/",
            "service_city",
            "Fence Company Naperville",
            "2026-09-25",
            "not-registered",
        ]
    )
    buffer = io.BytesIO()
    workbook.save(buffer)

    repository = ParentAwareRepository()
    service = make_service(repository)
    result = await service.import_pages(
        ImportPagesRequest(
            filename="pages.xlsx",
            xlsx_base64=base64.b64encode(buffer.getvalue()).decode("ascii"),
            dry_run=True,
        )
    )

    assert result["error_count"] == 0
    # The second row keeps the slug even though no registered page matches it.
    assert repository.created_payload is not None
    assert repository.created_payload["parent_campaign_page_id"] is None
    assert repository.created_payload["parent_slug"] == "not-registered"


async def test_update_page_only_sends_provided_fields() -> None:
    repository = FakeManagementRepository()
    repository.pages[188] = {"id": 188, "campaign_id": 1, "slug": "old-slug"}
    service = make_service(repository)

    result = await service.update_page(
        188,
        CampaignPageUpdateRequest(postal_code="60505", parent_slug="/fences/"),
    )

    assert result["ok"] is True
    assert repository.updated_page_payload == {
        "postal_code": "60505",
        "parent_slug": "fences",
    }


def test_update_page_rejects_an_empty_payload() -> None:
    with pytest.raises(ValueError):
        CampaignPageUpdateRequest().to_repository_payload()


def test_update_page_rejects_an_explicit_null_url() -> None:
    with pytest.raises(ValueError, match="url cannot be null"):
        CampaignPageUpdateRequest(url=None).to_repository_payload()


def test_youtube_channel_must_be_an_id_instead_of_a_url() -> None:
    with pytest.raises(ValueError):
        CampaignPageUpdateRequest(youtube_channel_id="https://youtube.com/channel/UCtest")


def test_create_page_accepts_published_for_an_already_optimized_page() -> None:
    request = CampaignPageCreateRequest(
        campaign_id=1,
        url="https://example.com/already-optimized/",
        primary_keyword="Already optimized",
        status="published",
    )

    assert request.status == "published"
    assert _page_setup_action(request.to_repository_payload()) == "skip"


async def test_create_post_group_normalizes_cities_and_state() -> None:
    repository = FakeManagementRepository()
    service = make_service(repository)

    result = await service.create_post_group(
        PostGroupCreateRequest(
            campaign_id=1,
            state="il",
            region_name="Chicago Metro",
            group_type="metro",
            cities=[
                {"city": " Chicago "},
                {"city": "chicago"},
                {"city": "Aurora", "state": "il"},
            ],
        )
    )

    payload = repository.post_group_payload
    assert payload is not None
    assert payload["state"] == "IL"
    assert payload["cities"] == [
        {"city": "Chicago", "state": "IL"},
        {"city": "Aurora", "state": "IL"},
    ]
    assert result["post_group"]["id"] == 501


def test_post_group_city_requires_a_state() -> None:
    with pytest.raises(ValueError):
        PostGroupUpdateRequest(cities=[{"city": "Chicago"}]).to_repository_payload()


def test_post_groups_endpoint_filters_and_uppercases_the_state() -> None:
    repository = FakeManagementRepository()
    repository.post_groups = [{"id": 501, "campaign_id": 1}]
    service = make_service(repository)
    app = FastAPI()
    app.include_router(router, prefix="/seo")
    app.dependency_overrides[seo_management_service] = lambda: service

    with TestClient(app) as client:
        response = client.get(
            "/seo/post-groups",
            params={
                "campaign_id": 1,
                "state": "il",
                "post_status": "defined",
                "search": "chicago",
            },
        )

    assert response.status_code == 200
    assert response.json() == [{"id": 501, "campaign_id": 1}]
    assert repository.post_group_filters["state"] == "IL"
    assert repository.post_group_filters["post_status"] == "defined"
    assert repository.post_group_filters["search"] == "chicago"


def test_update_page_endpoint_sends_only_the_edited_fields() -> None:
    repository = FakeManagementRepository()
    repository.pages[188] = {"id": 188, "campaign_id": 1}
    service = make_service(repository)
    app = FastAPI()
    app.include_router(router, prefix="/seo")
    app.dependency_overrides[seo_management_service] = lambda: service

    with TestClient(app) as client:
        response = client.put(
            "/seo/pages/188",
            json={"postal_code": "60505", "elementor_form_id": "a1b2c3d"},
        )

    assert response.status_code == 200
    assert repository.updated_page_payload == {
        "postal_code": "60505",
        "elementor_form_id": "a1b2c3d",
    }


def test_post_group_endpoints_create_update_and_delete() -> None:
    repository = FakeManagementRepository()
    service = make_service(repository)
    app = FastAPI()
    app.include_router(router, prefix="/seo")
    app.dependency_overrides[seo_management_service] = lambda: service

    with TestClient(app) as client:
        created = client.post(
            "/seo/post-groups",
            json={
                "campaign_id": 1,
                "state": "IL",
                "region_name": "Chicago Metro",
                "cities": [{"city": "Chicago"}],
            },
        )
        updated = client.put(
            "/seo/post-groups/501",
            json={
                "campaign_id": 2,
                "post_status": "defined",
                "post_title": "Nuevo título",
            },
        )
        deleted = client.delete("/seo/post-groups/501")

    assert created.status_code == 201
    assert created.json()["post_group"]["state"] == "IL"
    assert updated.status_code == 200
    # A partial update must not resend the untouched fields.
    assert repository.post_group_payload == {
        "campaign_id": 2,
        "post_status": "defined",
        "post_title": "Nuevo título",
    }
    assert deleted.status_code == 200


def test_import_template_endpoint_returns_a_workbook_the_importer_accepts() -> None:
    app = FastAPI()
    app.include_router(router, prefix="/seo")

    with TestClient(app) as client:
        response = client.get("/seo/pages/import-template")

    assert response.status_code == 200
    assert IMPORT_TEMPLATE_FILENAME in response.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(response.content))
    assert SHEET_NAME in workbook.sheetnames
    headers = [cell.value for cell in next(workbook[SHEET_NAME].iter_rows(max_row=1))]
    assert set(REQUIRED_COLUMNS).issubset(headers)
    for column in ("postal_code", "parent_slug", "elementor_form_id", "priority"):
        assert column in headers
    validations = {
        str(validation.sqref): validation.formula1
        for validation in workbook[SHEET_NAME].data_validations.dataValidation
    }
    assert validations["AH2:AH201"] == '"cities,services"'
    assert validations["AI2:AI201"] == '"pending,active,published"'


async def test_schedule_pages_persists_support_posts_mode_in_notes() -> None:
    repository = FakeManagementRepository()
    service = make_service(repository)

    result = await service.schedule_pages(
        SchedulePagesRequest(
            campaign_page_ids=[188],
            scheduled_for="2026-07-08",
            notes="Programada desde dashboard",
            support_posts={"mode": "state_defined"},
        )
    )

    assert result["inserted_count"] == 1
    assert repository.scheduled_for == date(2026, 7, 8)
    assert repository.schedule_notes == (
        "Programada desde dashboard [support_posts.mode=state_defined]"
    )


async def test_schedule_pages_only_queues_even_when_start_now_is_requested() -> None:
    repository = FakeManagementRepository()
    service = make_service(
        repository,
        page_execution_service=FakePageExecutionService(),
        submit_flow=FakeSubmitFlow(),
    )

    result = await service.schedule_pages(
        SchedulePagesRequest(
            campaign_page_ids=[188],
            scheduled_for=date.today().isoformat(),
            notes="Programada desde dashboard",
            start_now=True,
        )
    )

    assert result["inserted_count"] == 1
    assert result["started_count"] == 0
    assert result["started"] == []


async def test_run_daily_starts_pages_grouped_by_campaign_bot() -> None:
    class DueQueueRepository(FakeManagementRepository):
        async def get_due_queue(self, limit: int) -> list[dict[str, Any]]:
            return [
                {"id": 36, "campaign_page_id": 188, "campaign_id": 1, "notes": ""},
                {"id": 37, "campaign_page_id": 189, "campaign_id": 1, "notes": ""},
                {"id": 38, "campaign_page_id": 288, "campaign_id": 2, "notes": ""},
                {"id": 39, "campaign_page_id": 289, "campaign_id": 2, "notes": ""},
            ][:limit]

    repository = DueQueueRepository()
    repository.pages = {
        188: {
            "id": 188,
            "campaign_id": 1,
            "url": "https://example.com/campaign-one/page-a",
            "slug": "page-a",
            "campaign_name": "Campaign One",
            "primary_keyword": "Keyword A",
        },
        189: {
            "id": 189,
            "campaign_id": 1,
            "url": "https://example.com/campaign-one/page-b",
            "slug": "page-b",
            "campaign_name": "Campaign One",
            "primary_keyword": "Keyword B",
            "wp_page_id": 1002,
        },
        288: {
            "id": 288,
            "campaign_id": 2,
            "url": "https://example.com/campaign-two/page-a",
            "slug": "page-a",
            "campaign_name": "Campaign Two",
            "primary_keyword": "Keyword C",
            "wp_page_id": 2001,
            "has_successful_optimization": True,
        },
        289: {
            "id": 289,
            "campaign_id": 2,
            "url": "https://example.com/campaign-two/page-b",
            "slug": "page-b",
            "campaign_name": "Campaign Two",
            "primary_keyword": "Keyword D",
            "wp_page_id": 2002,
            "status": "published",
        },
    }
    bots = [
        Bot(name="Machine 1", bot_key="machine-1", bot_type="seo", capabilities=["seo.main"]),
        Bot(name="Machine 2", bot_key="machine-2", bot_type="seo", capabilities=["seo.main"]),
        Bot(name="Machine 3", bot_key="machine-3", bot_type="seo", capabilities=["seo.main"]),
    ]
    submit_flow = FakeSubmitFlow()
    service = make_service(
        repository,
        page_execution_service=FakePageExecutionService(),
        submit_flow=submit_flow,
        dispatcher=FakeDispatcher(bots),
    )

    result = await service.run_daily(RunDailyRequest(confirm=True, daily_limit=10))

    assert result["started_count"] == 4
    assert result["dispatched_count"] == 4
    assert result["pending_bot_count"] == 0
    assert result["already_started_count"] == 0
    requests_by_page = {request.campaign_page_id: request for request in submit_flow.requests}
    assert requests_by_page[188].preferred_bot_id == bots[0].id
    assert requests_by_page[189].preferred_bot_id == bots[0].id
    assert requests_by_page[288].preferred_bot_id == bots[1].id
    assert requests_by_page[289].preferred_bot_id == bots[1].id
    page_setups = {page_id: request.page_setup for page_id, request in requests_by_page.items()}
    assert all(page_setup is not None for page_setup in page_setups.values())
    assert page_setups[188] is not None
    assert page_setups[189] is not None
    assert page_setups[288] is not None
    assert page_setups[289] is not None
    assert page_setups[188].action == "create"
    assert page_setups[188].slug == "page-a"
    assert page_setups[189].action == "update"
    assert page_setups[189].wp_page_id == 1002
    assert page_setups[288].action == "skip"
    assert page_setups[289].action == "skip"


async def test_pending_seo_bot_marks_queue_running_without_duplicate_flow() -> None:
    class DueQueueRepository(FakeManagementRepository):
        async def get_due_queue(self, limit: int) -> list[dict[str, Any]]:
            return [{"id": 1734, "campaign_page_id": 872, "campaign_id": 1, "notes": ""}]

    class PendingSubmitFlow(FakeSubmitFlow):
        async def execute(self, request: Any) -> Any:
            response = await super().execute(request)
            response.dispatched_initial_execution.status = ExecutionStatus.PENDING
            return response

    repository = DueQueueRepository()
    repository.pages[872] = {
        "id": 872,
        "campaign_id": 1,
        "url": "https://example.com/commercial-fence-company-skokie-il/",
        "slug": "commercial-fence-company-skokie-il",
        "wp_page_id": 17778,
        "has_successful_optimization": True,
    }
    page_executions = FakePageExecutionService()
    submit_flow = PendingSubmitFlow()
    service = make_service(
        repository,
        page_execution_service=page_executions,
        submit_flow=submit_flow,
    )

    first = await service.run_daily(RunDailyRequest(confirm=True))
    second = await service.run_daily(RunDailyRequest(confirm=True))

    assert first["started"][0]["dispatch_status"] == "pending"
    assert second["started"][0]["dispatch_status"] == "already_started"
    assert first["pending_bot_count"] == 1
    assert second["already_started_count"] == 1
    assert repository.running_queue_ids == [1734, 1734]
    assert len(submit_flow.requests) == 1
    assert submit_flow.requests[0].page_setup.action == "skip"
    assert len(page_executions.linked) == 1


async def test_failed_initial_dispatch_does_not_leave_queue_running() -> None:
    class FailedSubmitFlow(FakeSubmitFlow):
        async def execute(self, request: Any) -> Any:
            response = await super().execute(request)
            response.dispatched_initial_execution.status = ExecutionStatus.FAILED
            response.dispatched_initial_execution.error_message = "Unsupported stage"
            return response

    repository = FakeManagementRepository()
    repository.pages[872] = {
        "id": 872,
        "campaign_id": 1,
        "url": "https://example.com/page/",
        "slug": "page",
        "has_successful_optimization": True,
    }
    service = make_service(
        repository,
        page_execution_service=FakePageExecutionService(),
        submit_flow=FailedSubmitFlow(),
    )

    result = await service._start_queue_flow(1734, 872)

    assert result["dispatch_status"] == "failed"
    assert repository.running_queue_ids == [1734]
    assert repository.queue_results == [
        {
            "queue_execution_id": 1734,
            "status": "failed",
            "result_status": "failed",
            "error_message": "Unsupported stage",
        }
    ]


async def test_claimed_queue_cannot_create_a_second_flow() -> None:
    repository = FakeManagementRepository()
    repository.pages[872] = {
        "id": 872,
        "campaign_id": 1,
        "url": "https://example.com/page/",
        "slug": "page",
        "has_successful_optimization": True,
    }
    repository.claimed_queue_ids.add(1734)
    submit_flow = FakeSubmitFlow()
    service = make_service(
        repository,
        page_execution_service=FakePageExecutionService(),
        submit_flow=submit_flow,
    )

    with pytest.raises(ValueError, match="already starting or no longer queued"):
        await service._start_queue_flow(1734, 872)

    assert submit_flow.requests == []


async def test_interrupted_flow_start_keeps_claim_and_records_error() -> None:
    class BrokenSubmitFlow(FakeSubmitFlow):
        async def execute(self, request: Any) -> Any:
            raise RuntimeError("database unavailable")

    repository = FakeManagementRepository()
    repository.pages[872] = {
        "id": 872,
        "campaign_id": 1,
        "url": "https://example.com/page/",
        "slug": "page",
        "has_successful_optimization": True,
    }
    service = make_service(
        repository,
        page_execution_service=FakePageExecutionService(),
        submit_flow=BrokenSubmitFlow(),
    )

    with pytest.raises(RuntimeError, match="database unavailable"):
        await service._start_queue_flow(1734, 872)

    assert repository.claimed_queue_ids == {1734}
    assert repository.start_errors == [
        (1734, "Flow startup interrupted: database unavailable")
    ]
